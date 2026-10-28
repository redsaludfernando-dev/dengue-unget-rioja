# -*- coding: utf-8 -*-
"""
Actualiza los datos y las fotos de la web "Prevención del dengue — UNGET Rioja".

Lee:
  - WEB/BASE_WEB.xlsx ................. registro de actividades, evidencias, ajustes y configuración
  - PROMOCION DE LA SALUD/ACTIVIDADES DENGUE 2026 - PROMOCIÓN DE LA SALUD.xlsx
  - CONTROL VECTORIAL/Formato Control*.xlsx, Formato Vigilancia*.xlsx, Formato Nebuliz*.xlsx
  - PROMOCION DE LA SALUD/PUBLICACIONES EN REDES SOCIALES.txt
Escribe:
  - WEB/sitio/data/datos.js y datos.json
  - WEB/sitio/img/... (fotos reducidas, sin metadatos ni coordenadas)

Uso:  python3 actualizar.py            (desde cualquier carpeta)
"""
import os, sys, re, json, glob, unicodedata, datetime as dt
import pandas as pd
from openpyxl import load_workbook
from PIL import Image, ImageOps

AQUI = os.path.dirname(os.path.abspath(__file__))
WEB = os.path.dirname(AQUI)
REPORTES = os.path.dirname(WEB)
SITIO = os.path.join(WEB, 'sitio')
BASE = os.path.join(WEB, 'BASE_WEB.xlsx')
AVISOS = []

def aviso(msg):
    AVISOS.append(msg)

def sin_tildes(s):
    s = unicodedata.normalize('NFKD', str(s))
    return ''.join(c for c in s if not unicodedata.combining(c)).upper()

def buscar(patron):
    r = sorted(glob.glob(os.path.join(REPORTES, patron)))
    if not r:
        raise SystemExit(f'No encontré el archivo: {patron}')
    return r[0]

# ------------------------------------------------------------------ distritos
DISTRITOS = [
    ('rioja', 'Rioja', ''),
    ('elias-soplin-vargas', 'Elías Soplín Vargas', 'Segunda Jerusalén'),
    ('nueva-cajamarca', 'Nueva Cajamarca', ''),
    ('pardo-miguel', 'Pardo Miguel', 'Naranjos'),
    ('awajun', 'Awajún', 'Bajo Naranjillo'),
    ('yuracyacu', 'Yuracyacu', ''),
    ('san-fernando', 'San Fernando', ''),
    ('posic', 'Pósic', ''),
    ('yorongos', 'Yorongos', ''),
    ('moyobamba', 'Moyobamba', 'Valle de la Conquista y Atumplaya'),
]
REGLAS = [  # el orden importa: la primera coincidencia gana
    ('TODA LA JURISDICCION', 'todos'),
    ('BAJO NARANJILLO', 'awajun'), ('AWAJUN', 'awajun'),
    ('CONQUISTA', 'moyobamba'), ('ATUMPLAYA', 'moyobamba'), ('MOYOBAMBA', 'moyobamba'),
    ('POSIC', 'posic'), ('YORONGOS', 'yorongos'), ('SAN FERNANDO', 'san-fernando'), ('YURACYACU', 'yuracyacu'),
    ('ELIAS SOPLIN', 'elias-soplin-vargas'), ('SEGUNDA JERUSALEN', 'elias-soplin-vargas'), ('AMAZONAS', 'elias-soplin-vargas'),
    ('PARDO MIGUEL', 'pardo-miguel'), ('NARANJOS', 'pardo-miguel'), ('PIONEROS', 'pardo-miguel'), ('SANTA ROSA DEL MIRADOR', 'pardo-miguel'),
    ('AGUAS VERDES', 'pardo-miguel'), ('AGUAS CLARAS', 'pardo-miguel'),
    ('NUEVA CAJAMARCA', 'nueva-cajamarca'), ('NUEVO CAJAMARCA', 'nueva-cajamarca'), ('NARANJILLO', 'nueva-cajamarca'), ('UCRANIA', 'nueva-cajamarca'),
    ('SORITOR', 'nueva-cajamarca'), ('SJR', 'nueva-cajamarca'), ('LA VICTORIA', 'nueva-cajamarca'), ('LA UNION', 'nueva-cajamarca'), ('SAN JUAN BAUTISTA', 'nueva-cajamarca'),
    ('SAN FRANCISCO', 'awajun'), ('PORVENIR', 'rioja'), ('CAPIRONAL', 'rioja'), ('RIOJA', 'rioja'),
]
def distrito_id(*textos):
    t = sin_tildes(' '.join(str(x) for x in textos if x is not None and str(x) != 'nan'))
    for clave, did in REGLAS:
        if clave in t:
            return did
    return None

# Establecimiento (EESS_Localidad de los formatos) -> (localidad, distrito)
EESS = [
    ('BAJO NARANJILLO (CC.PP. SAN FRANCISCO)', 'San Francisco', 'awajun'),
    ('BAJO NARANJILLO', 'Bajo Naranjillo', 'awajun'),
    ('LA VICTORIA', 'La Victoria Nueva', 'nueva-cajamarca'),
    ('SORITOR', 'San Juan de Río Soritor', 'nueva-cajamarca'),
    ('NUEVA RIOJA', 'Rioja', 'rioja'),
    ('PORVENIR', 'El Porvenir', 'rioja'),
    ('SEGUNDA JERUSALEN', 'Segunda Jerusalén', 'elias-soplin-vargas'),
    ('AGUAS VERDES', 'Aguas Verdes', 'pardo-miguel'),
    ('AGUAS CLARAS', 'Aguas Claras', 'pardo-miguel'),
    ('NARANJOS', 'Naranjos', 'pardo-miguel'),
    ('NARANJILLO', 'Naranjillo', 'nueva-cajamarca'),
    ('NUEVA CAJAMARCA', 'Nueva Cajamarca', 'nueva-cajamarca'),
    ('UCRANIA', 'Ucrania', 'nueva-cajamarca'),
    ('LA UNION', 'La Unión', 'nueva-cajamarca'),
    ('SAN FERNANDO', 'San Fernando', 'san-fernando'),
    ('POSIC', 'Pósic', 'posic'),
    ('YORONGOS', 'Yorongos', 'yorongos'),
]
def eess(nombre):
    t = sin_tildes(nombre)
    for clave, loc, did in EESS:
        if sin_tildes(clave) in t:
            return loc, did
    aviso(f'Establecimiento sin distrito reconocido: {nombre}')
    return str(nombre).title(), None

MESES = {'ENERO': 1, 'FEBRERO': 2, 'MARZO': 3, 'ABRIL': 4, 'MAYO': 5, 'JUNIO': 6, 'JULIO': 7, 'AGOSTO': 8,
         'SETIEMBRE': 9, 'SEPTIEMBRE': 9, 'OCTUBRE': 10, 'NOVIEMBRE': 11, 'DICIEMBRE': 12}
def mes_num(texto, fecha=None):
    t = sin_tildes(texto).strip() if texto is not None and str(texto) != 'nan' else ''
    if t in MESES:
        return MESES[t]
    if fecha is not None and not pd.isna(fecha):
        return pd.Timestamp(fecha).month
    return None

def num(x):
    try:
        if x is None or (isinstance(x, float) and pd.isna(x)) or str(x).strip() in ('', '-'):
            return None
        return float(x)
    except Exception:
        return None

def ent(x):
    v = num(x)
    return int(round(v)) if v is not None else 0

def fecha_iso(x):
    if x is None or (isinstance(x, float) and pd.isna(x)) or str(x).strip() == '':
        return ''
    if isinstance(x, (dt.date, dt.datetime, pd.Timestamp)):
        return pd.Timestamp(x).strftime('%Y-%m-%d')
    s = str(x).strip()
    m = re.match(r'(\d{4})-(\d{2})(?:-(\d{2}))?', s)
    if m:
        return s[:10] if m.group(3) else s[:7]
    m = re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})', s)
    if m:
        return f'{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}'
    aviso(f'Fecha no reconocida: {s}')
    return ''

# ------------------------------------------------------------------ fotos
CACHE_F = os.path.join(SITIO, 'img', 'cache.json')
try:
    CACHE = json.load(open(CACHE_F, encoding='utf-8'))
except Exception:
    CACHE = {}

def caja_recorte(recorte):
    """Devuelve (x0, y0, x1, y1) en fracciones. Formatos:
       'abajo' (quita el 24 % inferior), 'abajo:0.4', 'arriba:0.2', 'caja:x0,y0,x1,y1'."""
    r = (recorte or '').strip().lower()
    if not r:
        return (0, 0, 1, 1)
    try:
        if r == 'abajo':
            return (0, 0, 1, 0.76)
        if r.startswith('abajo:'):
            return (0, 0, 1, 1 - float(r.split(':')[1]))
        if r.startswith('arriba:'):
            return (0, float(r.split(':')[1]), 1, 1)
        if r.startswith('caja:'):
            x0, y0, x1, y1 = [float(v) for v in r.split(':')[1].split(',')]
            return (x0, y0, x1, y1)
    except Exception:
        pass
    aviso(f'Recorte no reconocido: "{recorte}" (se usa la foto completa)')
    return (0, 0, 1, 1)

def procesar_foto(origen, destino, lado, calidad, recorte=''):
    firma = f'{os.path.getmtime(origen):.0f}|{recorte}|{lado}'
    rel = os.path.relpath(destino, SITIO)
    if os.path.exists(destino) and CACHE.get(rel, {}).get('firma') == firma:
        return tuple(CACHE[rel]['tam'])
    with Image.open(origen) as im:
        im = ImageOps.exif_transpose(im).convert('RGB')   # gira según la cámara
        w, h = im.size
        x0, y0, x1, y1 = caja_recorte(recorte)             # quita sellos de fecha y coordenadas
        if (x0, y0, x1, y1) != (0, 0, 1, 1):
            im = im.crop((int(w * x0), int(h * y0), int(w * x1), int(h * y1)))
        im.thumbnail((lado, lado), Image.LANCZOS)
        os.makedirs(os.path.dirname(destino), exist_ok=True)
        im.save(destino, 'JPEG', quality=calidad, optimize=True, progressive=True)  # sin EXIF ni GPS
        CACHE[rel] = {'firma': firma, 'tam': list(im.size)}
        return im.size

def fotos_de(carpeta, archivos, recortes, prefijo):
    salida = []
    if not carpeta or not archivos:
        return salida
    lista = [a.strip() for a in str(archivos).split(';') if a.strip()]
    rec = [r.strip() for r in str(recortes or '').split(';')]
    for i, nombre in enumerate(lista):
        origen = os.path.join(REPORTES, carpeta, nombre)
        if not os.path.exists(origen):
            aviso(f'Foto no encontrada: {carpeta}/{nombre}')
            continue
        r = rec[i] if i < len(rec) else (rec[0] if rec else '')
        base = f'{prefijo}-{i + 1}'
        w, h = procesar_foto(origen, os.path.join(SITIO, 'img', base + '.jpg'), 1280, 80, r)
        procesar_foto(origen, os.path.join(SITIO, 'img', base + '-t.jpg'), 560, 72, r)
        salida.append({'src': f'img/{base}.jpg', 'th': f'img/{base}-t.jpg', 'w': w, 'h': h})
    return salida

# ------------------------------------------------------------------ base web
def leer_hoja(wb, nombre):
    ws = wb[nombre]
    filas = list(ws.iter_rows(values_only=True))
    cab = [str(c).strip() if c is not None else '' for c in filas[0]]
    out = []
    for f in filas[1:]:
        if f is None or all(v is None or str(v).strip() == '' for v in f):
            continue
        out.append({cab[i]: f[i] for i in range(len(cab)) if cab[i]})
    return out

def col(fila, *nombres):
    for n in nombres:
        for k, v in fila.items():
            if sin_tildes(k).startswith(sin_tildes(n)):
                return v
    return None

wb = load_workbook(BASE, data_only=True)
config = {sin_tildes(col(f, 'PARAMETRO')): col(f, 'VALOR') for f in leer_hoja(wb, 'CONFIG')}
corte = fecha_iso(config.get('FECHA DE CORTE')) or dt.date.today().isoformat()
proxima = None
if config.get('PROXIMA ACTIVIDAD'):
    p = str(config['PROXIMA ACTIVIDAD']).split('|')
    if len(p) >= 3:
        proxima = {'fecha': fecha_iso(p[0]), 'd': distrito_id(p[1]), 'texto': p[2].strip()}

actividades = []
for f in leer_hoja(wb, 'ACTIVIDADES'):
    aid = str(col(f, 'ID') or '').strip()
    if not aid or not col(f, 'ACTIVIDAD'):
        continue
    d = distrito_id(col(f, 'DISTRITO'))
    if d is None:
        aviso(f'{aid}: distrito no reconocido "{col(f, "DISTRITO")}"')
    publicar = sin_tildes(col(f, 'PUBLICAR') or 'SI').strip() != 'NO'
    fotos = fotos_de(col(f, 'CARPETA'), col(f, 'FOTOS'), col(f, 'RECORTE'), f'a/{aid}') if publicar else []
    aliados = [a.strip() for a in str(col(f, 'ALIADOS') or '').split(';') if a.strip()]
    actividades.append({
        'id': aid, 'fecha': fecha_iso(col(f, 'FECHA')), 'comp': str(col(f, 'COMPONENTE') or '').strip(),
        'd': d or 'todos', 'lugar': str(col(f, 'LUGAR') or '').strip(), 'tipo': str(col(f, 'TIPO') or '').strip(),
        'titulo': str(col(f, 'ACTIVIDAD')).strip(), 'desc': str(col(f, 'DESCRIPCI') or '').strip(),
        'part': str(col(f, 'PARTICIPANTES') or '').strip(), 'res': str(col(f, 'RESULTADO') or '').strip(),
        'aliados': aliados, 'fotos': fotos,
    })

evidencias = []
for f in leer_hoja(wb, 'EVIDENCIAS CV'):
    eid = str(col(f, 'ID') or '').strip()
    if not eid or sin_tildes(col(f, 'PUBLICAR') or 'SI').strip() == 'NO':
        continue
    fotos = fotos_de(col(f, 'CARPETA'), col(f, 'FOTOS'), col(f, 'RECORTE'), f'e/{eid}')
    if fotos:
        evidencias.append({'id': eid, 'mes': fecha_iso(col(f, 'MES')), 'comp': str(col(f, 'COMPONENTE') or '').strip(),
                           'd': distrito_id(col(f, 'DISTRITO')) or 'todos', 'desc': str(col(f, 'DESCRIPCI') or '').strip(),
                           'fotos': fotos})

ajustes = leer_hoja(wb, 'AJUSTES') if 'AJUSTES' in wb.sheetnames else []

# ------------------------------------------------------------------ promoción (consolidado)
PROMSA = buscar('PROMOCION DE LA SALUD/ACTIVIDADES DENGUE*PROMOCI*.xlsx')
wp = load_workbook(PROMSA, data_only=True)

def hoja_que_empieza(nombre):
    for s in wp.sheetnames:
        if sin_tildes(s).strip().startswith(sin_tildes(nombre)):
            return wp[s]
    raise SystemExit(f'No encontré la hoja {nombre} en {os.path.basename(PROMSA)}')

# Recojos: lista LUGAR | FECHA | CANTIDAD RECOGIDA TN (columna BC:BE)
ws = hoja_que_empieza('RECOJOS')
recojos, programados = [], []
cab = None
for fila in ws.iter_rows():
    for c in fila:
        if isinstance(c.value, str) and sin_tildes(c.value).strip() == 'LUGAR':
            cab = (c.row, c.column)
            break
    if cab:
        break
if not cab:
    raise SystemExit('No encontré la tabla LUGAR / FECHA / CANTIDAD en la hoja RECOJOS')
r = cab[0] + 1
while r <= ws.max_row:
    lugar = ws.cell(r, cab[1]).value
    fecha = ws.cell(r, cab[1] + 1).value
    cant = ws.cell(r, cab[1] + 2).value
    if lugar is None and fecha is None:
        r += 1
        if r > cab[0] + 200:
            break
        continue
    if isinstance(lugar, str) and 'TOTAL' in sin_tildes(lugar):
        break
    lugar_txt = re.sub(r'^\d+\s+', '', str(lugar).strip())
    f_iso = fecha_iso(fecha)
    t = num(cant)
    for a in ajustes:
        if sin_tildes(col(a, 'HOJA') or '') == 'RECOJOS' and sin_tildes(col(a, 'LUGAR') or '') in sin_tildes(lugar_txt) \
                and fecha_iso(col(a, 'FECHA')) == f_iso and num(col(a, 'VALOR EN')) == t:
            aviso(f'Ajuste aplicado: recojo {lugar_txt} {f_iso}: {t} -> {col(a, "VALOR A USAR")} t')
            t = num(col(a, 'VALOR A USAR'))
    d = distrito_id(lugar_txt)
    if d is None:
        aviso(f'Recojo con lugar no reconocido: {lugar_txt}')
    item = {'d': d, 'lugar': lugar_txt.title().replace(' De ', ' de ').replace(' La ', ' la '), 'fecha': f_iso}
    if t is None:
        programados.append(item)
    else:
        item['t'] = round(t, 3)
        recojos.append(item)
    r += 1

TILDES = {'CHAVE': 'Chávez', 'CHAVEZ': 'Chávez', 'AGUILA': 'Águila', 'VELASQUEZ': 'Velásquez', 'LOPEZ': 'López',
          'HERNANDEZ': 'Hernández', 'BILINGUE': 'Bilingüe', 'RIO': 'Río', 'JERUSALEN': 'Jerusalén'}
def nombre_ie(texto):
    pal = str(texto).replace('}', '').split()
    out = []
    for i, w in enumerate(pal):
        W = sin_tildes(w)
        if W == 'IE':
            out.append('IE')
        elif W in TILDES:
            out.append(TILDES[W])
        elif w.isdigit():
            out.append(w)
        elif i > 0 and W in ('DE', 'DEL', 'Y'):
            out.append(w.lower())
        else:
            out.append(w[:1].upper() + w[1:].lower())
    n = ' '.join(out)
    return n if n.startswith('IE ') else 'IE ' + n

# Concurso "Al cole seguro"
ws = hoja_que_empieza('IE CONCURSO')
concurso, dist_actual = [], None
for fila in ws.iter_rows(min_row=1, values_only=True):
    vals = list(fila)
    if len(vals) < 7:
        continue
    c_dist, c_ie, c_fecha, c_kg, c_gan = vals[2], vals[3], vals[4], vals[5], vals[6]
    if isinstance(c_dist, str) and sin_tildes(c_dist).strip() not in ('DISTRITOS', '') and 'TOTAL' not in sin_tildes(c_dist):
        dist_actual = c_dist
    if c_ie is None or sin_tildes(c_ie).startswith('IE PARTICIPANTES'):
        continue
    if not isinstance(c_fecha, (dt.datetime, dt.date)):
        continue
    estado = sin_tildes(c_gan or '').strip()
    concurso.append({'d': distrito_id(c_ie, dist_actual), 'ie': nombre_ie(c_ie), 'fecha': fecha_iso(c_fecha),
                     'kg': num(c_kg) or 0, 'ganadora': estado == 'GANADOR', 'participo': 'NO PARTICIPO' not in estado})

# Ordenanzas
ws = hoja_que_empieza('ORDENANZAS')
ordenanzas = []
for fila in ws.iter_rows(values_only=True):
    vals = [v for v in fila]
    if len(vals) >= 4 and vals[1] and vals[2] and sin_tildes(vals[1]).strip() != 'DISTRITO':
        ordenanzas.append({'d': distrito_id(vals[1]), 'ordenanza': str(vals[2]).strip(), 'periodo': str(vals[3] or '').strip()})

# Comités de vigilancia comunitaria (solo conteos; no se exportan nombres, DNI ni teléfonos)
ws = hoja_que_empieza('COMITE DE VIGILANCIA')
comites, actual, periodo = [], None, ''
for fila in ws.iter_rows(values_only=True):
    a, b = fila[0], fila[1] if len(fila) > 1 else None
    if isinstance(a, str):
        ta = sin_tildes(a).strip()
        if 'COMITE DE VIGILANCIA' in ta and ':' in a:
            lugar = a.split(':', 1)[1].strip()
            actual = {'lugar': lugar, 'd': distrito_id(lugar), 'periodo': periodo, 'integrantes': 0}
            comites.append(actual)
            continue
        if re.fullmatch(r'[A-Z\- ]+', ta) and any(m in ta for m in MESES):
            periodo = a.strip().title().replace('- ', '–').replace('-', '–')
            actual = None
            continue
    if actual is not None and isinstance(a, (int, float)) and b not in (None, ''):
        actual['integrantes'] += 1
comites = [c for c in comites if c['lugar'] and c['integrantes'] > 0]

# Capacitaciones
ws = hoja_que_empieza('CAPACITACIONES A DOCENTES')
docentes = []
for fila in ws.iter_rows(min_row=3, values_only=True):
    v = list(fila)
    if len(v) >= 8 and isinstance(v[2], (int, float)) and v[3]:
        docentes.append({'d': distrito_id(v[3], v[5]), 'ie': nombre_ie(v[3]), 'fecha': fecha_iso(v[6]), 'n': ent(v[7])})

ws = hoja_que_empieza('CAPACITACION A L')
lideres, mes_actual = [], None
for fila in ws.iter_rows(min_row=3, values_only=True):
    v = list(fila)
    if len(v) < 5:
        continue
    if isinstance(v[1], str) and sin_tildes(v[1]).strip() in MESES:
        mes_actual = MESES[sin_tildes(v[1]).strip()]
    if v[2] and v[3] is not None and num(v[3]):
        lideres.append({'d': distrito_id(v[2]), 'lugar': str(v[2]).strip().title(), 'mes': mes_actual, 'n': ent(v[3]),
                        'publico': str(v[4] or '').strip()})

ws = hoja_que_empieza('CAPACITACION A AC')
agentes = []
for fila in ws.iter_rows(min_row=3, values_only=True):
    v = list(fila)
    if len(v) >= 5 and v[1] and v[2] and num(v[3]):
        f = v[4]
        if isinstance(f, (dt.datetime, dt.date)) and f.year < 2026:
            aviso(f'Capacitación a agentes/actores con año {f.year}; se tomó como 2026')
            f = f.replace(year=2026)
        agentes.append({'d': distrito_id(v[2]), 'publico': str(v[1]).strip(), 'fecha': fecha_iso(f), 'n': ent(v[3])})

def publico(texto):
    t = sin_tildes(texto)
    if 'ESTUDIANTE' in t: return 'Estudiantes'
    if 'DOCENTE' in t: return 'Docentes'
    if 'PADRE' in t: return 'Padres de familia'
    if 'PERSONAL DE SALUD' in t: return 'Personal de salud'
    if 'PENSION' in t: return 'Adultos mayores (Pensión 65)'
    if 'AGENTE' in t: return 'Agentes comunitarios'
    if 'ACTOR' in t: return 'Actores sociales'
    return 'Líderes comunitarios'
for x in lideres: x['grupo'] = publico(x['publico'])
for x in agentes: x['grupo'] = publico(x['publico'])
for x in docentes: x['grupo'] = 'Docentes'

# Publicaciones en redes
publicaciones = []
try:
    with open(buscar('PROMOCION DE LA SALUD/PUBLICACIONES EN REDES*.txt'), encoding='utf-8', errors='replace') as fh:
        publicaciones = [l.strip() for l in fh if l.strip().startswith('http')]
except SystemExit:
    pass

# Afiches de difusión
flyers = []
carpeta_fly = 'PROMOCION DE LA SALUD/FLYERS PARA RECOJO DE CRIADEROS'
lista_fly = config.get('AFICHES') or ''
if os.path.isdir(os.path.join(REPORTES, carpeta_fly)):
    nombres = [n.strip() for n in str(lista_fly).split(';') if n.strip()] or \
        sorted(n for n in os.listdir(os.path.join(REPORTES, carpeta_fly)) if n.lower().endswith(('.jpg', '.jpeg', '.png')))[:8]
    flyers = fotos_de(carpeta_fly, ';'.join(nombres), '', 'd/afiche')

# ------------------------------------------------------------------ control vectorial
def leer_formato(patron, fila_cab):
    df = pd.read_excel(buscar(patron), header=fila_cab)
    df.columns = [str(c).strip() for c in df.columns]
    return df

DEP = ['Tanque alto', 'Tanque bajo', 'Barril o cilindro', 'Sansón o bidón', 'Baldes, bateas o tinajas',
       'Llantas', 'Floreros y maceteros', 'Latas y botellas', 'Inservibles', 'Otros']

dfc = leer_formato('CONTROL VECTORIAL/Formato Control*.xlsx', 7)
dfc = dfc[dfc['EESS_Localidad'].notna()]
control = []
for _, f in dfc.iterrows():
    loc, d = eess(f['EESS_Localidad'])
    control.append({
        'd': d, 'loc': loc, 'mes': mes_num(f.get('mes'), f.get('fec_ini')), 'ini': fecha_iso(f.get('fec_ini')), 'fin': fecha_iso(f.get('fec_ter')),
        'vprog': ent(f.get('viv_prog')), 'vi': ent(f.get('viv_insp')), 'vp': ent(f.get('viv_pos')),
        'vc': ent(f.get('viv_cer')), 'vr': ent(f.get('viv_ren')), 'vd': ent(f.get('viv_desh')),
        'ri': ent(f.get('rec_insp')), 'rp': ent(f.get('rec_pos')), 'rtq': ent(f.get('rec_TQ')), 'rtf': ent(f.get('rec_TF')),
        'rdes': ent(f.get('ec_dest')), 'larv': round(num(f.get('larvicida')) or 0, 1),
        'dep': [ent(f.get(f'{i}_pos')) for i in range(1, 11)],
    })

dfv = leer_formato('CONTROL VECTORIAL/Formato Vigilancia*.xlsx', 7)
dfv = dfv[dfv['EESS_Localidad'].notna()]
vigilancia = []
for _, f in dfv.iterrows():
    loc, d = eess(f['EESS_Localidad'])
    vi, vp, ri, rp = ent(f.get('viv_insp')), ent(f.get('viv_pos')), ent(f.get('rec_insp')), ent(f.get('rec_pos'))
    vigilancia.append({
        'd': d, 'loc': loc, 'mes': mes_num(f.get('mes'), f.get('fec_ini')), 'fecha': fecha_iso(f.get('fec_ini')),
        'vi': vi, 'vp': vp, 'ri': ri, 'rp': rp,
        'ia': round(100 * vp / vi, 2) if vi else None, 'ir': round(100 * rp / ri, 2) if ri else None,
        'ib': round(100 * rp / vi, 2) if vi else None,
        'dep': [ent(f.get(f'{i}_pos')) for i in range(1, 11)],
    })

dfn = leer_formato('CONTROL VECTORIAL/Formato Nebuliz*.xlsx', 4)
for c in ['Provincia', 'Distrito', 'UNGETS', 'Localidad']:
    if c in dfn.columns:
        dfn[c] = dfn[c].ffill()
dfn = dfn[dfn['zonas'].notna()]
cercos = {}
for _, f in dfn.iterrows():
    loc, d = eess(f['Localidad'])
    zona = re.sub(r'\s+', ' ', str(f['zonas'])).strip()
    clave = (loc, zona)
    if clave not in cercos:
        cercos[clave] = {'d': d, 'loc': loc, 'zona': zona, 'mes': mes_num(f.get('mes'), f.get('inicio')), 'vueltas': [],
                         'mz': ent(f.get('mz_pro')), 'insect': 0.0}
    c = cercos[clave]
    c['insect'] += num(f.get('insec_litros')) or 0
    if num(f.get('viv_neb')) is None:
        continue                                  # vuelta sin datos de viviendas (pendiente o vacía)
    c['vueltas'].append({'v': str(f.get('Vuelta') or '').strip(), 'fecha': fecha_iso(f.get('inicio')),
                         'prog': ent(f.get('viv_pro')), 'neb': ent(f.get('viv_neb')), 'cer': ent(f.get('viv_cer')),
                         'ren': ent(f.get('viv_ren')), 'des': ent(f.get('viv_desh'))})
    c['mz'] = c['mz'] or ent(f.get('mz_pro'))
neb = [c for c in cercos.values() if c['vueltas']]
# Quita cercos repetidos: misma localidad y misma primera vuelta (fecha y viviendas nebulizadas)
neb.sort(key=lambda c: (-len(c['vueltas']), -c['mz']))
vistos, nebulizacion = set(), []
for c in neb:
    v1 = c['vueltas'][0]
    firma = (c['loc'], v1['fecha'], v1['neb'])
    if firma in vistos:
        aviso(f'Cerco de nebulización repetido omitido: {c["loc"]} — {c["zona"]} ({v1["fecha"]})')
        continue
    vistos.add(firma)
    c['insect'] = round(c['insect'], 1)
    nebulizacion.append(c)
nebulizacion.sort(key=lambda c: (c['vueltas'][0]['fecha'], c['loc']))

# ------------------------------------------------------------------ salida
datos = {
    'meta': {'corte': corte, 'periodo': str(config.get('PERIODO') or '').strip(), 'entidad': str(config.get('ENTIDAD') or 'UNGET Rioja'),
             'generado': dt.datetime.now(dt.timezone(dt.timedelta(hours=-5))).strftime('%d/%m/%Y a las %H:%M'), 'proxima': proxima},
    'distritos': [{'id': i, 'nombre': n, 'nota': o} for i, n, o in DISTRITOS],
    'depositos': DEP,
    'vigilancia': vigilancia, 'control': control, 'nebulizacion': nebulizacion,
    'recojos': recojos, 'recojosProgramados': programados, 'concurso': concurso, 'ordenanzas': ordenanzas,
    'comites': comites, 'capacitaciones': {'docentes': docentes, 'lideres': lideres, 'agentes': agentes},
    'actividades': actividades, 'evidencias': evidencias,
    'difusion': {'publicaciones': publicaciones, 'afiches': flyers},
}
os.makedirs(os.path.join(SITIO, 'data'), exist_ok=True)
with open(CACHE_F, 'w', encoding='utf-8') as fh:
    json.dump(CACHE, fh)
texto = json.dumps(datos, ensure_ascii=False, separators=(',', ':'))
with open(os.path.join(SITIO, 'data', 'datos.json'), 'w', encoding='utf-8') as fh:
    fh.write(texto)
with open(os.path.join(SITIO, 'data', 'datos.js'), 'w', encoding='utf-8') as fh:
    fh.write('window.DATOS = ' + texto + ';\n')

print(f'Corte: {corte}')
print(f'Actividades: {len(actividades)} (con foto: {sum(1 for a in actividades if a["fotos"])}) | Evidencias CV: {len(evidencias)}')
print(f'Vigilancia: {len(vigilancia)} muestreos | Control larvario: {len(control)} intervenciones | Nebulización: {len(nebulizacion)} cercos')
print(f'Recojos: {len(recojos)} campañas, {sum(r["t"] for r in recojos):.2f} t | Programados: {len(programados)}')
print(f'Concurso IE: {len(concurso)} | Ordenanzas: {len(ordenanzas)} | Comités: {len(comites)}')
print(f'Capacitaciones: docentes {sum(x["n"] for x in docentes)}, líderes y otros {sum(x["n"] for x in lideres)}, agentes y actores {sum(x["n"] for x in agentes)}')
if AVISOS:
    print('\nAVISOS:')
    for a in AVISOS:
        print(' -', a)
